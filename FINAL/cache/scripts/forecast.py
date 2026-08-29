# -*- coding: utf-8 -*-
"""
forecast.py — 基准情景 × 多轮递归预测,输出逐月预测明细(预测到 2031 年底)。

预测序列(每小区 3 条,各自独立训练):
  综合均价(主)、精装二手、以及 毛坯二手(北控)/ 别墅二手(帝泊湾,按表内可用类型自适应)。
每条序列输出:基准 P50 + 80% 区间(P10/P90)。

事件驱动:未来 66 个月事件流 = 延续最近 12 个月政策/市场事件节奏(见 features.future_event_flow),
事件特征(ev_net6/12/24)参与模型输入;事件弹性 β 用于诊断事件方向(见 models.estimate_event_beta)。

表格输出目录:ROOT/output/  (独立于代码与原始数据)
"""
import os

import pandas as pd

from config_loader import CFG
from features import build_events, load_extra_events
from models import recursive_forecast, summarize_rounds, estimate_event_beta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(ROOT, "output")

# ---- 从集中配置读取预测参数 ----
_PC = CFG["predict"]
FORECAST_MONTHS = _PC["forecast_months"]
N_ROUNDS = _PC["n_rounds"]


def series_for_target(old, target):
    """从二手房表抽取某条预测序列(统一为 date/price/is_actual)。
    剔除该列无价格的行(NaN),避免把历史空行误当预测目标。"""
    col = {"综合": "price", "精装二手": "jingzhuang", "毛坯二手": "maopi",
           "别墅二手": "bieshu"}[target]
    s = old[["date", "is_actual"]].copy()
    s["price"] = old[col]
    s = s.dropna(subset=["price"])
    return s.sort_values("date").reset_index(drop=True)


def targets_for(old):
    """该小区需要预测的序列:综合 + 精装二手 + (毛坯二手 或 别墅二手)。"""
    ts = ["综合", "精装二手"]
    if old["bieshu"].notna().any():
        ts.append("别墅二手")
    elif old["maopi"].notna().any():
        ts.append("毛坯二手")
    return ts


def forecast_complex(complex_dict, events_series, n=FORECAST_MONTHS,
                     rounds=N_ROUNDS, seed=2026, on_round=None):
    """
    对一个小区:对每条序列独立做 多轮(bootstrap)递归预测,
    输出 {序列名: DataFrame(date, P10, P50, P90, mean)} + _meta。
    on_round: 可选回调 on_round(已完成轮数, 总轮数),用于进度显示。
    """
    old = complex_dict["old"]
    targets = targets_for(old)
    result = {}
    counter = [0]
    total = rounds * len(targets)

    def _cb(done, total_rounds):
        counter[0] += 1
        if on_round is not None:
            on_round(counter[0], total)

    for t in targets:
        series = series_for_target(old, t)
        rounds_df = recursive_forecast(series, events_series, "基准", n=n,
                                       rounds=rounds, seed=seed, on_round=_cb)
        result[t] = summarize_rounds(rounds_df)
    result["_meta"] = {
        "targets": targets,
        "beta": {t: estimate_event_beta(series_for_target(old, t), events_series)[0]
                 for t in targets},
    }
    return result


def _merged_df(result):
    """各序列合并为一张表:date + 每序列 P10/P50/P90/mean(整数)。"""
    frames = []
    for t in result["_meta"]["targets"]:
        df = result[t].copy()
        df.columns = ["date"] + [f"{t}_{c}" for c in df.columns[1:]]
        frames.append(df)
    out = frames[0]
    for f in frames[1:]:
        out = out.merge(f, on="date")
    num_cols = [c for c in out.columns if c != "date"]
    out[num_cols] = out[num_cols].round(0).astype(int)
    return out


def export(name, result):
    """输出预测表(XLSX 模板样式)到 output/ 目录。"""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out = _merged_df(result)
    # 日期由预测结果直接构造(YYYY-MM 月末)
    dates = pd.to_datetime(out["date"] + "-01", format="%Y-%m-%d") + pd.offsets.MonthEnd(0)
    return export_xlsx(name, out, dates, result)


def export_xlsx(name, out, dates, result=None):
    """
    按模板样式输出 XLSX(模板: predictions_新力帝泊湾_价格数据表.xlsx):
      字体=等线(表头12号/数据22号), 行高=表头15.5/数据27.5,
      A列宽=15.58, 水平垂直居中, 细边框, 日期列格式 d-mmm。
    dates: 月末日期数组(与 out 行数一致)。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    targets = [c[:-4] for c in out.columns if c.endswith("_P10")]
    cols = ["date"] + [f"{t}_{q}" for t in targets for q in ("P10", "P50", "P90", "mean")]

    wb = Workbook()
    ws = wb.active
    ws.title = f"predictions_{name}"[:31]

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # 表头:等线 12 号,行高 15.5
    for j, h in enumerate(cols, 1):
        cell = ws.cell(1, j, h)
        cell.font = Font(name="等线", size=12)
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 15.5

    # 数据:等线 22 号,行高 27.5,整数;日期列月末日期
    for i in range(len(out)):
        r = i + 2
        ws.cell(r, 1, dates[i].to_pydatetime()).number_format = "d-mmm"
        for j, col in enumerate(cols[1:], 2):
            ws.cell(r, j, int(out[col].iloc[i]))
        for j in range(1, len(cols) + 1):
            cell = ws.cell(r, j)
            cell.font = Font(name="等线", size=22)
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[r].height = 27.5

    # 列宽:A 列 15.58(与模板一致);其余列加宽,防中文表头截断(如 精装二手_P90)
    ws.column_dimensions["A"].width = 15.58
    from openpyxl.utils import get_column_letter
    for col in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # 三情景 sheet(基准/乐观/悲观 P50)——result['_scenarios'] 存在时写入
    scen = result.get("_scenarios")
    if scen:
        ws2 = wb.create_sheet("情景")
        thin2 = Side(style="thin")
        border2 = Border(left=thin2, right=thin2, top=thin2, bottom=thin2)
        center2 = Alignment(horizontal="center", vertical="center")
        first_t = result["_meta"]["targets"][0]
        dates2 = list(scen[first_t]["基准"].keys())
        ws2.cell(1, 1, "月份")
        for j, sn in enumerate(("基准", "乐观", "悲观"), 2):
            ws2.cell(1, j, f"{sn}P50")
        for j in range(1, 4):
            ws2.cell(1, j).font = Font(name="等线", size=12, bold=True)
            ws2.cell(1, j).alignment = center2
            ws2.cell(1, j).border = border2
        ws2.row_dimensions[1].height = 15.5
        for i, d in enumerate(dates2, 2):
            ws2.cell(i, 1, d)
            for j, sn in enumerate(("基准", "乐观", "悲观"), 2):
                v = round(float(scen[first_t][sn][d]), 0)
                ws2.cell(i, j, v)
            for j in range(1, 4):
                ws2.cell(i, j).font = Font(name="等线", size=12)
                ws2.cell(i, j).alignment = center2
                ws2.cell(i, j).border = border2
            ws2.row_dimensions[i].height = 27.5
        ws2.column_dimensions["A"].width = 15.58
        ws2.column_dimensions["B"].width = 16
        ws2.column_dimensions["C"].width = 16
        ws2.column_dimensions["D"].width = 16

    xlsx_path = os.path.join(OUTPUT_DIR, f"predictions_{name}.xlsx")
    wb.save(xlsx_path)
    return xlsx_path


def print_overview(name, old, result):
    beta = result["_meta"]["beta"]
    print(f"\n[{name}] 预测至 2031-12 · 当前(2026.06)综合均价 {old['price'].iloc[-1]:,.0f} 元/㎡")
    print(f"  事件弹性 β: " + "  ".join(f"{t}={b:+.4f}" for t, b in beta.items()))
    for t in result["_meta"]["targets"]:
        df = result[t]
        last = df.iloc[-1]
        chg = last["P50"] / df.iloc[0]["P50"] - 1
        print(f"  {t}: 2026.07 起 {df.iloc[0]['P50']:,.0f} → 2031.12 {last['P50']:,.0f} "
              f"({chg:+.1%})  80%区间 [{last['P10']:,.0f}, {last['P90']:,.0f}]")


if __name__ == "__main__":
    from load_data import discover, load_complex
    complexes = [load_complex(fp) for fp in discover()]
    all_events = pd.concat([c["events"] for c in complexes if c["events"] is not None],
                           ignore_index=True)
    events = build_events(all_events, load_extra_events())
    for c in complexes:
        result = forecast_complex(c, events)
        path = export(c["name"], result)
        print_overview(c["name"], c["old"], result)
        print("  预测明细 →", path)
